import os
import PyPDF2
import re
import docx

# Get the first PDF file in the "pdf" folder
pdf_folder = "pdf"
pdf_file_name = next(f for f in os.listdir(pdf_folder) if f.endswith('.pdf'))
pdf_path = os.path.join(pdf_folder, pdf_file_name)

with open(pdf_path, 'rb') as pdf_file:
    pdf_reader = PyPDF2.PdfReader(pdf_file)
    page = pdf_reader.pages[0]
    content = page.extract_text()

    # Extract name and gender
    pattern = r"(\d{3}\s+)([\u4e00-\u9fa5]+)\s+(男|女)"
    match = re.search(pattern, content)
    
    if match:
        name = match.group(2)
        gender = match.group(3)
        print(f"Name: {name}, Gender: {gender}")  # Debug extracted name and gender
    else:
        print("No match found")  # Debug if pattern doesn't match

# Load the Word document template
doc = docx.Document('output.docx')

# Replace the placeholders with the extracted data
if name and gender:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "姓名：" in cell.text:
                    cell.text = cell.text.replace("姓名：", f"姓名：{name}")
                    print(f"Set name in cell: {cell.text}")  # Debug name in cell
                if "性别：" in cell.text:
                    cell.text = cell.text.replace("性别：", f"性别：{gender}")
                    print(f"Set gender in cell: {cell.text}")  # Debug gender in cell

    # Save the modified Word document
    doc.save('output.docx')
    print("Saved output.docx")  # Debug when the document is saved
else:
    print("Name and/or gender not found")  # Debug if name and/or gender not found

